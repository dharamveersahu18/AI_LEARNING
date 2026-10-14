
from groq import Groq
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from pathlib import Path
from pypdf import PdfReader
from docx import Document
import os
import json
import time

# Load .env from project root
# Load environment variables from project root
env_path = Path(__file__).resolve().parent / ".env"
load_dotenv(env_path)

# Check the key without printing the secret
api_key = os.getenv("GROQ_API_KEY")

if not api_key:
    raise ValueError(f"GROQ_API_KEY not found. Check: {env_path}")

# Create Groq client
client = Groq(api_key=api_key)

model = "openai/gpt-oss-20b"
# Job description
class JobD(BaseModel):
    role: str
    required_skills: list[str]
    preferred_skills: list[str]
    minimum_experience: float | None
    education_requirements: list[str]
    responsibility: list[str]


jobD_schema = JobD.model_json_schema()


# Put your actual job description text here

# Job description
job_description = """
Job Title: Junior Python Developer

Required Skills:
- Python programming
- Object-oriented programming
- SQL and database fundamentals
- REST APIs
- Git and GitHub

Preferred Skills:
- FastAPI or Flask
- Basic knowledge of AI and LLMs
- Experience with GitHub projects

Minimum Experience: 0 to 1 year

Education Requirements:
- Bachelor's degree in Computer Science, IT, or a related field

Responsibilities:
- Develop and maintain Python applications
- Build and test REST APIs
- Work with databases
- Debug and document code
- Collaborate with the development team
"""
if not job_description.strip():
    raise ValueError(
        "Job description is empty. Please add a job description first."
    )

system_prompt = f"""
You are an expert HR assistant.
Your job is to analyze a job description and extract structured information from it.

Return ONLY valid JSON matching this schema:
{json.dumps(jobD_schema)}

IMPORTANT:
- Do NOT return the schema itself.
- Do not return fields like "properties", "title", or "type".
- Fill the schema with actual information extracted from the job description.
- If minimum experience is not mentioned, return null.
- If information for a list is missing, return an empty list.
- Do not invent information.
"""

user_prompt = f"""
Analyze the following job description:
{job_description}
"""

message_system = {
    "role": "system",
    "content": system_prompt
}

message_user = {
    "role": "user",
    "content": user_prompt
}

response_format = {
    "type": "json_object"
}

messages = [message_system, message_user]

# Send request to Groq
response = client.chat.completions.create(
    model=model,
    messages=messages,
    response_format=response_format

)

# Get the AI response
answer = response.choices[0].message.content

# Convert JSON string into a Python dictionary
raw_json = answer
job_data = json.loads(raw_json)
job = JobD(**job_data)

print(job.minimum_experience)
print(job.education_requirements)


# Parse resume and matching result
class MatchResult(BaseModel):
    score: float
    details: dict


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = Field(default_factory=list)


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    total_experience_years: float | None = None
    skills: list[str] = Field(default_factory=list)
    experince: list[Experience] = Field(default_factory=list)
    education: list[str] = Field(default_factory=list)
    projects: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


resume_schema = Resume.model_json_schema()


def final_score(job, resume):
    match_schema = MatchResult.model_json_schema()

    prompt = f"""
You are an HR recruiter.
Compare the candidate's resume with the job description.

Job Description:
{job.model_dump_json(indent=2)}

Candidate Resume:
{resume.model_dump_json(indent=2)}

Return JSON matching this schema:
{json.dumps(match_schema)}

Include these details in the "details" field:
1. Candidate name
2. Matching skills
3. Missing important skills
4. Whether the required experience is met
5. A short final verdict

Set "score" to the overall match percentage from 0 to 100.
Keep the response concise and easy to read.
Do not invent information.
"""

    message = {
        "role": "user",
        "content": prompt
    }

    messages = [message]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format
    )

    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)


def parse_resume(resume_text):
    system_prompt = f"""
You are an expert resume parser.
Extract information from the resume based on its meaning,
not only based on exact section headings.

Different resumes may use different headings, for example:
- Experience
- Professional Experience
- Work History
- Employment
- Internships

These may all contain relevant experience.
Skills may also appear in the skills section, work experience,
internships, or projects.

Return ONLY valid JSON matching this schema:
{json.dumps(resume_schema)}

Important rules:
1. Do not invent information.
2. If a value is not available, return null.
3. If a list has no information, return an empty list.
4. Include internships inside experience.
5. Extract skills mentioned across the entire resume.
"""

    user_prompt = f"""
Parse the following resume:
{resume_text}
"""

    message_system = {
        "role": "system",
        "content": system_prompt
    }

    message_user = {
        "role": "user",
        "content": user_prompt
    }

    messages = [message_system, message_user]

    response_format = {
        "type": "json_object"
    }

    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format=response_format,
            max_completion_tokens=4096
    )

    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)

    return resume


def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"

    return text


def read_docs(file_path):
    document = Document(file_path)
    text = ""

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"

    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docs(file_path)
    else:
        return None


resume_folder = Path("resumes")
all_results = []

for file_path in resume_folder.iterdir():
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue

    print("\nProcessing:", file_path.name)

    resume_text = read_resume(file_path)
    if not resume_text or not resume_text.strip():
        print("No readable text found. Skipping:", file_path.name)
        continue

    parsed_resume = parse_resume(resume_text)

    time.sleep(5)

    result = final_score(job, parsed_resume)

    time.sleep(5)

    print("Score:", result.score)

    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })

all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)

top_2 = all_results[:2]
worst_2 = all_results[-2:][::-1]

print("Top 2 candidates")
for candidate in top_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])

print("Lowest 2 candidates")
for candidate in worst_2:
    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])
